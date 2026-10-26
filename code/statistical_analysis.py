from pathlib import Path
from statistics import median

from openpyxl import load_workbook, Workbook
from scipy.stats import wilcoxon, rankdata

models = [
    "gpt-6-astra",
    "llama-4-maverick-instruct",
    "gemini-4-argon",
    "claude-opus-5-5",
    "qwen3.8-max-2026-09-02",
    "deepseek-v4-pro"
]

# Folder paths
PROJECT_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_DIR / "data"
RESULTS_DIR = PROJECT_DIR / "results"

results = []

for model in models:

    # Load data
    wb = load_workbook(DATA_DIR / f"results_{model}.xlsx")
    ws = wb["equitability"]

    cyrillic = []
    latin = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            cyrillic.append(row[1])
            latin.append(row[2])

    wb.close()

    # Wilcoxon signed-rank test
    test = wilcoxon(cyrillic, latin)

    # Effect size (rank-biserial correlation)
    differences = []

    for c, l in zip(cyrillic, latin):
        if c != l:
            differences.append(c - l)

    ranks = rankdata([abs(d) for d in differences])

    positive = 0
    negative = 0

    for i in range(len(differences)):
        if differences[i] > 0:
            positive += ranks[i]
        else:
            negative += ranks[i]

    r = (positive - negative) / (positive + negative)

    results.append([
        model,
        median(cyrillic),
        median(latin),
        test.statistic,
        test.pvalue,
        r
    ])

# Holm correction
p_values = [row[4] for row in results]
order = sorted(range(len(p_values)), key=lambda i: p_values[i])

adjusted = [0] * len(p_values)
previous = 0

for k, i in enumerate(order):
    p = min(1, (len(p_values) - k) * p_values[i])
    previous = max(previous, p)
    adjusted[i] = previous

# Save results to Excel
wb = Workbook()
ws = wb.active

ws.append([
    "Model", "Median Cyrillic", "Median Latin",
    "W", "p", "p (Holm)", "r"
])

for i in range(len(results)):
    model, med_cyr, med_lat, w, p, r = results[i]

    ws.append([
        model,
        round(float(med_cyr), 4),
        round(float(med_lat), 4),
        round(float(w), 1),
        round(float(p), 3),
        round(float(adjusted[i]), 3),
        round(float(r), 2)
    ])

RESULTS_DIR.mkdir(exist_ok=True)
wb.save(RESULTS_DIR / "table.xlsx")

print("Results saved to results/table.xlsx")